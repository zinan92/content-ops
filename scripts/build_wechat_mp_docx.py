#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit(
        "Missing python-docx. Run this with the bundled Codex Python runtime or install python-docx."
    ) from exc

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover - optional rendering helper
    Image = None
    ImageDraw = None
    ImageFont = None


IMAGE_RE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_image(markdown_path: Path, image_ref: str) -> Path | None:
    raw = image_ref.strip()
    if raw.startswith("file://"):
        raw = raw[7:]
    path = Path(raw)
    if path.is_absolute():
        return path if path.exists() else None
    candidate = markdown_path.parent / raw
    return candidate if candidate.exists() else None


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(12)


def font(size: int, bold: bool = False) -> Any:
    if ImageFont is None:
        return None
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            try:
                return ImageFont.truetype(str(path), size=size, index=1 if bold and path.suffix == ".ttc" else 0)
            except Exception:
                continue
    return ImageFont.load_default()


def build_article_spine_image(draft: dict[str, Any]) -> Path | None:
    if Image is None or ImageDraw is None:
        return rendered_image(draft, "article-spine.png")
    media_dir = Path(str(draft.get("rendered_package_dir") or ""))
    if not media_dir:
        return None
    media_dir.mkdir(parents=True, exist_ok=True)
    output_path = media_dir / "article-spine.png"
    image = Image.new("RGB", (1840, 1040), "#f7f5ee")
    draw = ImageDraw.Draw(image)
    green = "#14845f"
    dark = "#15231d"
    muted = "#51605a"
    border = "#d5ded8"
    draw.text((96, 82), "文章主线", fill=dark, font=font(66, bold=True))
    cards = [
        ("01", "先定义赌注", "知道自己在赌什么"),
        ("02", "再找不变变量", "避开无效机会"),
        ("03", "最后加速迭代", "把 AI 放进系统"),
    ]
    x_positions = [96, 668, 1240]
    for x, (num, title, subtitle) in zip(x_positions, cards, strict=False):
        draw.rounded_rectangle((x, 250, x + 500, 720), radius=26, fill="#ffffff", outline=border, width=3)
        draw.text((x + 48, 315), num, fill=green, font=font(42, bold=True))
        draw.text((x + 48, 415), title, fill=dark, font=font(50, bold=True))
        draw.text((x + 48, 520), subtitle, fill=muted, font=font(38))
    draw.rectangle((96, 785, 1744, 910), fill="#e8f3ee")
    draw.rectangle((96, 785, 106, 910), fill=green)
    draw.text((140, 825), "先定义赌注，再识别不变变量，最后用 AI 加速迭代。", fill=dark, font=font(42, bold=True))
    image.save(output_path)
    return output_path


def rendered_image(draft: dict[str, Any], filename: str) -> Path | None:
    for raw in draft.get("rendered_images") or []:
        path = Path(str(raw))
        if path.name == filename and path.exists():
            return path
    media_dir = Path(str(draft.get("rendered_package_dir") or ""))
    if media_dir:
        candidate = media_dir / filename
        if candidate.exists():
            return candidate
    return None


def add_image_with_caption(document: Document, image_path: Path | None, caption: str) -> None:
    if not image_path:
        return
    document.add_picture(str(image_path), width=Inches(6.2))
    paragraph = document.add_paragraph()
    paragraph.alignment = 1
    run = paragraph.add_run(caption)
    run.font.size = Pt(9)
    run.italic = True


def add_markdown_line(document: Document, markdown_path: Path, line: str) -> None:
    stripped = line.strip()
    if not stripped:
        document.add_paragraph()
        return

    image_match = IMAGE_RE.fullmatch(stripped)
    if image_match:
        image_path = resolve_image(markdown_path, image_match.group(1))
        if image_path:
            document.add_picture(str(image_path), width=Inches(6.2))
        return

    if stripped.startswith("# "):
        document.add_heading(stripped[2:].strip(), level=1)
        return
    if stripped.startswith("## "):
        document.add_heading(stripped[3:].strip(), level=2)
        return
    if stripped.startswith("### "):
        document.add_heading(stripped[4:].strip(), level=3)
        return
    if stripped.startswith("- "):
        document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return

    add_paragraph(document, stripped)


def build_docx(draft_json_path: Path, output_path: Path | None = None) -> Path:
    draft = read_json(draft_json_path)
    markdown_path = Path(draft.get("markdown_path") or draft_json_path.with_suffix(".md"))
    if not markdown_path.exists():
        raise FileNotFoundError(f"Markdown draft not found: {markdown_path}")

    output_path = output_path or draft_json_path.with_suffix(".docx")
    document = Document()

    styles = document.styles
    styles["Normal"].font.name = "PingFang SC"
    styles["Normal"].font.size = Pt(12)

    title = str(draft.get("title") or "").strip()
    digest = str(draft.get("digest") or "").strip()
    if title:
        document.add_heading(title, level=0)
    if digest:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(digest)
        run.italic = True
        run.font.size = Pt(11)

    add_image_with_caption(
        document,
        build_article_spine_image(draft),
        "文章主线：先定义赌注，再识别不变变量，最后用 AI 加速迭代。",
    )
    add_image_with_caption(
        document,
        rendered_image(draft, "decision-checklist.png"),
        "判断先于行动：先定义赌注、失效信号和退出条件。",
    )

    for line in markdown_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("---") or line.startswith("platform:") or line.startswith("source_content_id:") or line.startswith("local_id:"):
            continue
        if title and line.strip() == f"# {title}":
            continue
        if digest and line.strip() == f"> {digest}":
            continue
        add_markdown_line(document, markdown_path, line)

    add_image_with_caption(
        document,
        rendered_image(draft, "framework-layers.png"),
        "不要只追应用层热闹，真正长期有效的是底层框架。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a WeChat MP importable DOCX from an outbox draft JSON.")
    parser.add_argument("draft_json", type=Path, help="Path to outbox/drafts/wechat_mp/*.json")
    parser.add_argument("--output", type=Path, default=None, help="Optional output .docx path")
    args = parser.parse_args()

    output = build_docx(args.draft_json, args.output)
    print(output)


if __name__ == "__main__":
    main()
